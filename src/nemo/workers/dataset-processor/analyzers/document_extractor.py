"""
Document text extraction for PII analysis.

Extracts plain text from PDF, DOCX, and PPTX files so they can be passed to
the Presidio-based ``analyze_text()`` function.  Handles common failure modes:

- **Password-protected files**: returns ``None`` with a warning log so the
  caller can skip PII analysis gracefully rather than crashing.
- **Oversized files**: files larger than ``MAX_DOCUMENT_SIZE_BYTES`` are
  rejected before any parsing to prevent OOM.

The extracted text is returned as a single string.  The caller is responsible
for further truncation before passing to Presidio (``analyze_text()`` already
enforces a 5 MB limit on the text side).

Supported formats
-----------------
* PDF  – via ``pymupdf`` (PyMuPDF / ``fitz``).  Only text layers are read;
         scanned image-only PDFs with no embedded text return an empty string.
* DOCX – via ``python-docx``.  Paragraphs and table cells are extracted.
* PPTX – via ``python-pptx``.  Text frames from all slide shapes are extracted.

Usage
-----
    from analyzers.document_extractor import extract_document_text

    text = extract_document_text("/path/to/file.pdf", ".pdf")
    if text is None:
        # file is password-protected or extraction failed fatally
        ...
    elif text == "":
        # file parsed successfully but contained no extractable text
        ...
    else:
        result = analyze_text(text)
"""

import os
from observability_client_runtime import get_logger
from typing import Optional

logger = get_logger()

# Maximum binary file size to attempt parsing (100 MB).
# Avoids unbounded memory use when parsing unusually large documents.
MAX_DOCUMENT_SIZE_BYTES = 100 * 1024 * 1024  # 100 MB


def _check_file_size(path: str) -> bool:
    """Return True if the file is within the allowed size limit.

    Logs a warning and returns False when the file exceeds
    ``MAX_DOCUMENT_SIZE_BYTES`` so the caller can skip analysis.
    """
    try:
        size = os.path.getsize(path)
    except OSError as exc:
        logger.warning("Could not determine file size for %s: %s", path, exc)
        return False

    if size > MAX_DOCUMENT_SIZE_BYTES:
        logger.warning(
            "Document %s (%d bytes) exceeds the %d-byte limit for PII extraction; skipping.",
            path,
            size,
            MAX_DOCUMENT_SIZE_BYTES,
        )
        return False
    return True


# ---------------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------------

def _extract_pdf(path: str) -> Optional[str]:
    """Extract text from a PDF using PyMuPDF (fitz).

    Returns:
        Extracted text string (may be empty for image-only PDFs).
        ``None`` if the PDF is password-protected or cannot be opened.
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        logger.error("pymupdf is not installed; cannot extract text from PDF files.")
        return None

    try:
        doc = fitz.open(path)
    except fitz.EmptyFileError:
        logger.warning("PDF file is empty: %s", path)
        return ""
    except Exception as exc:
        logger.warning("Failed to open PDF %s: %s", path, exc)
        return None

    # Encrypted / password-protected PDFs expose needs_pass = True.
    # fitz.open() succeeds but page iteration fails without authentication.
    if doc.needs_pass:
        logger.warning(
            "PDF %s is password-protected; skipping PII analysis.", path
        )
        doc.close()
        return None

    parts = []
    try:
        for page_num, page in enumerate(doc):
            try:
                parts.append(page.get_text())
            except Exception as exc:
                logger.warning(
                    "Failed to extract text from page %d of %s: %s", page_num, path, exc
                )
    finally:
        doc.close()
    return "\n".join(parts)


def _extract_docx(path: str) -> Optional[str]:
    """Extract text from a DOCX file using python-docx.

    Extracts paragraph text and table cell text.

    Returns:
        Extracted text string.
        ``None`` if the file is password-protected or structurally invalid.
    """
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError:
        logger.error("python-docx is not installed; cannot extract text from DOCX files.")
        return None

    try:
        doc = Document(path)
    except PackageNotFoundError:
        logger.warning("DOCX %s is not a valid Open XML package (possibly encrypted).", path)
        return None
    except Exception as exc:
        error_msg = str(exc).lower()
        if "encrypt" in error_msg or "password" in error_msg or "protected" in error_msg:
            logger.warning(
                "DOCX %s appears to be password-protected; skipping PII analysis.", path
            )
            return None
        logger.warning("Failed to open DOCX %s: %s", path, exc)
        return None

    parts = [para.text for para in doc.paragraphs if para.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts)


def _extract_pptx(path: str) -> Optional[str]:
    """Extract text from a PPTX file using python-pptx.

    Extracts text from all text frames across all slides.

    Returns:
        Extracted text string.
        ``None`` if the file is password-protected or structurally invalid.
    """
    try:
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError
    except ImportError:
        logger.error("python-pptx is not installed; cannot extract text from PPTX files.")
        return None

    try:
        prs = Presentation(path)
    except PackageNotFoundError:
        logger.warning("PPTX %s is not a valid Open XML package (possibly encrypted).", path)
        return None
    except Exception as exc:
        error_msg = str(exc).lower()
        if "encrypt" in error_msg or "password" in error_msg or "protected" in error_msg:
            logger.warning(
                "PPTX %s appears to be password-protected; skipping PII analysis.", path
            )
            return None
        logger.warning("Failed to open PPTX %s: %s", path, exc)
        return None

    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = para.text
                if text:
                    parts.append(text)

    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Public dispatcher
# ---------------------------------------------------------------------------

_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
}


def extract_document_text(path: str, extension: str) -> Optional[str]:
    """Extract plain text from a document file for PII analysis.

    Dispatches to the appropriate format-specific extractor based on
    *extension*.  Pre-checks the file size to avoid OOM when parsing
    unusually large documents.

    Args:
        path:      Absolute filesystem path to the document.
        extension: Lowercase file extension including the leading dot
                   (e.g. ``".pdf"``).

    Returns:
        Extracted text as a string (may be empty if the file has no text
        layer).  Returns ``None`` when:

        - The file exceeds ``MAX_DOCUMENT_SIZE_BYTES``.
        - The file is password-protected.
        - The file cannot be parsed due to a fatal error.
        - The extension is not supported.
    """
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        logger.debug("No document extractor for extension %r; skipping.", extension)
        return None

    if not _check_file_size(path):
        return None

    logger.info("Extracting text from %s document: %s", extension, path)
    try:
        text = extractor(path)
    except Exception as exc:
        logger.warning("Unexpected error extracting text from %s: %s", path, exc)
        return None

    if text is None:
        logger.info("Text extraction returned None for %s (likely protected or invalid).", path)
    elif not text.strip():
        logger.info("Text extraction returned empty content for %s.", path)

    return text
