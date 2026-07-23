"""Additional unit tests for data_sources/unstructured.py filling gaps left by
test_unstructured_extraction.py: connect()/get_documents() orchestration,
_process_file cleanup, timeout branches, ImportError fallbacks for optional
extraction libraries, exception paths, and RTF edge cases.
"""

from pathlib import Path
import sys
from unittest import mock

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from data_sources.unstructured import UnstructuredDataSource


def _make_source(tmp_path: Path, **kwargs) -> UnstructuredDataSource:
    return UnstructuredDataSource(dataset_id="dataset", temp_dir=tmp_path, **kwargs)


class TestConnect:
    def test_connect_with_file_keys_filter_skips_listing(self, tmp_path, monkeypatch):
        source = _make_source(tmp_path / "work", file_keys=["a.txt", "b.txt"])
        with mock.patch.object(source, "_list_dataset_files") as mock_list:
            source.connect()
        mock_list.assert_not_called()
        assert source._files == ["a.txt", "b.txt"]
        assert (tmp_path / "work").is_dir()

    def test_connect_without_file_keys_filter_lists_files(self, tmp_path, monkeypatch):
        mount = tmp_path / "store"
        data_dir = mount / "datasets/dataset/data_files"
        data_dir.mkdir(parents=True)
        (data_dir / "a.txt").write_text("hi", encoding="utf-8")
        monkeypatch.setenv("NEMO_DEFAULT_STORE_ROOT", str(mount))

        source = _make_source(tmp_path / "work")
        source.connect()
        assert source._files == ["datasets/dataset/data_files/a.txt"]


class TestListDatasetFilesSkipsHiddenAndResultFiles:
    def test_skips_dotfiles_and_result_files(self, tmp_path, monkeypatch):
        mount = tmp_path / "store"
        data_dir = mount / "datasets/dataset/data_files"
        data_dir.mkdir(parents=True)
        (data_dir / ".hidden.txt").write_text("x", encoding="utf-8")
        (data_dir / "job_result.json").write_text("{}", encoding="utf-8")
        (data_dir / "keep.txt").write_text("keep me", encoding="utf-8")
        monkeypatch.setenv("NEMO_DEFAULT_STORE_ROOT", str(mount))

        source = _make_source(tmp_path / "work")
        files = source._list_dataset_files()
        assert files == ["datasets/dataset/data_files/keep.txt"]

    def test_missing_base_dir_returns_empty_list(self, tmp_path, monkeypatch):
        monkeypatch.setenv("NEMO_DEFAULT_STORE_ROOT", str(tmp_path / "does-not-exist"))
        source = _make_source(tmp_path / "work")
        assert source._list_dataset_files() == []


class TestGetDocuments:
    def test_yields_documents_and_skips_files_with_no_text(self, tmp_path):
        source = _make_source(tmp_path)
        source._files = ["a.txt", "b.txt", "c.txt"]

        docs_by_key = {
            "a.txt": mock.Mock(content="hello"),
            "b.txt": None,
        }

        def fake_process(file_key):
            if file_key == "c.txt":
                raise RuntimeError("boom")
            return docs_by_key[file_key]

        with mock.patch.object(source, "_process_file", side_effect=fake_process):
            docs = list(source.get_documents())

        assert len(docs) == 1
        assert docs[0].content == "hello"

    def test_empty_file_list_does_not_raise_division_by_zero(self, tmp_path):
        source = _make_source(tmp_path)
        source._files = []
        assert list(source.get_documents()) == []


class TestProcessFileCleanup:
    def test_local_temp_file_removed_after_extraction(self, tmp_path, monkeypatch):
        source = _make_source(tmp_path / "work")
        source.prefix = "datasets/dataset/data_files/"
        file_key = "datasets/dataset/data_files/a.txt"

        real_src = tmp_path / "a.txt"
        real_src.write_text("content", encoding="utf-8")
        monkeypatch.setattr("utils.data_store.posix_path", lambda key: real_src)

        doc = source._process_file(file_key)
        assert doc is not None
        local_path = (tmp_path / "work") / "a.txt"
        assert not local_path.exists()

    def test_local_temp_file_removed_even_when_extraction_raises(self, tmp_path, monkeypatch):
        source = _make_source(tmp_path / "work")
        source.prefix = "datasets/dataset/data_files/"
        file_key = "datasets/dataset/data_files/a.txt"

        real_src = tmp_path / "a.txt"
        real_src.write_text("content", encoding="utf-8")
        monkeypatch.setattr("utils.data_store.posix_path", lambda key: real_src)

        with mock.patch.object(source, "_extract_text", side_effect=RuntimeError("boom")):
            with pytest.raises(RuntimeError):
                source._process_file(file_key)

        local_path = (tmp_path / "work") / "a.txt"
        assert not local_path.exists()

    def test_no_text_extracted_returns_none(self, tmp_path, monkeypatch):
        source = _make_source(tmp_path / "work")
        source.prefix = "datasets/dataset/data_files/"
        file_key = "datasets/dataset/data_files/empty.txt"

        real_src = tmp_path / "empty.txt"
        real_src.write_text("", encoding="utf-8")
        monkeypatch.setattr("utils.data_store.posix_path", lambda key: real_src)

        assert source._process_file(file_key) is None


class TestIsFileTooLarge:
    def test_large_file_is_skipped(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "big.txt"
        path.write_bytes(b"x")
        with mock.patch.object(Path, "stat") as mock_stat:
            mock_stat.return_value.st_size = source.MAX_PARSE_FILE_SIZE_BYTES + 1
            assert source._is_file_too_large(path) is True

    def test_small_file_is_not_skipped(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "small.txt"
        path.write_text("x", encoding="utf-8")
        assert source._is_file_too_large(path) is False

    def test_stat_oserror_treated_as_not_too_large(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "gone.txt"
        with mock.patch.object(Path, "stat", side_effect=OSError("gone")):
            assert source._is_file_too_large(path) is False

    def test_oversized_file_skips_extraction_entirely(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "big.txt"
        path.write_text("some content", encoding="utf-8")
        with mock.patch.object(source, "_is_file_too_large", return_value=True):
            assert source._extract_text(path) == ""


class TestExtractTextFormats:
    def test_tsv_extraction(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "data.tsv"
        path.write_text("col1\tcol2\nval1\tval2\n", encoding="utf-8")
        text = source._extract_text(path)
        assert "col1 col2" in text
        assert "val1 val2" in text

    def test_valid_json_is_pretty_printed(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "data.json"
        path.write_text('{"a": 1, "b": [1, 2]}', encoding="utf-8")
        text = source._extract_text(path)
        assert '"a": 1' in text

    def test_malformed_json_returns_raw_content(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "data.jsonl"
        path.write_text('{"a": 1}\n{"b": 2}\n', encoding="utf-8")
        text = source._extract_text(path)
        assert '{"a": 1}' in text

    def test_html_extraction_reads_raw_markup(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "page.html"
        path.write_text("<html><body>Hello</body></html>", encoding="utf-8")
        text = source._extract_text(path)
        assert "Hello" in text

    def test_yaml_and_ini_extraction(self, tmp_path):
        source = _make_source(tmp_path)
        yaml_path = tmp_path / "config.yaml"
        yaml_path.write_text("key: value\n", encoding="utf-8")
        assert "key: value" in source._extract_text(yaml_path)

        ini_path = tmp_path / "config.ini"
        ini_path.write_text("[section]\nkey=value\n", encoding="utf-8")
        assert "key=value" in source._extract_text(ini_path)

    def test_unknown_extension_falls_back_to_text_read(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "data.weird"
        path.write_text("plain fallback text", encoding="utf-8")
        assert "plain fallback text" in source._extract_text(path)

    def test_unknown_extension_unreadable_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "adir.weird"
        path.mkdir()  # a directory can't be opened as text -> IsADirectoryError
        assert source._extract_text(path) == ""

    def test_generic_exception_during_extraction_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.txt"
        path.write_text("hi", encoding="utf-8")
        with mock.patch("builtins.open", side_effect=RuntimeError("disk error")):
            assert source._extract_text(path) == ""


class TestGuessSuffixFromMime:
    def test_docx_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"not a real docx but has extension guess")
        with mock.patch("mimetypes.guess_type", return_value=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", None
        )):
            assert source._guess_suffix_from_mime(path) == ".docx"

    def test_legacy_msword_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"legacy word doc bytes")
        with mock.patch("mimetypes.guess_type", return_value=("application/msword", None)):
            assert source._guess_suffix_from_mime(path) == ".docx"

    def test_odt_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"not a real odt")
        with mock.patch("mimetypes.guess_type", return_value=(
            "application/vnd.oasis.opendocument.text", None
        )):
            assert source._guess_suffix_from_mime(path) == ".odt"

    def test_xls_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"legacy excel bytes")
        with mock.patch("mimetypes.guess_type", return_value=("application/vnd.ms-excel", None)):
            assert source._guess_suffix_from_mime(path) == ".xls"

    def test_rtf_mime_type_variants(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"not really rtf")
        with mock.patch("mimetypes.guess_type", return_value=("application/rtf", None)):
            assert source._guess_suffix_from_mime(path) == ".rtf"
        with mock.patch("mimetypes.guess_type", return_value=("text/rtf", None)):
            assert source._guess_suffix_from_mime(path) == ".rtf"

    def test_no_match_returns_empty_string(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"totally unknown content")
        with mock.patch("mimetypes.guess_type", return_value=(None, None)):
            assert source._guess_suffix_from_mime(path) == ""

    def test_pdf_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"not actually a pdf by signature")
        with mock.patch("mimetypes.guess_type", return_value=("application/pdf", None)):
            assert source._guess_suffix_from_mime(path) == ".pdf"

    def test_pptx_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"not actually a pptx by signature")
        with mock.patch("mimetypes.guess_type", return_value=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation", None
        )):
            assert source._guess_suffix_from_mime(path) == ".pptx"

    def test_xlsx_mime_type(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        path.write_bytes(b"not actually an xlsx by signature")
        with mock.patch("mimetypes.guess_type", return_value=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None
        )):
            assert source._guess_suffix_from_mime(path) == ".xlsx"


class TestSniffBinaryType:
    def test_unreadable_file_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "gone.bin"
        with mock.patch("builtins.open", side_effect=OSError("no access")):
            assert source._sniff_binary_type(path) == ""

    def test_corrupt_zip_container_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "corrupt.bin"
        path.write_bytes(b"PK\x03\x04" + b"\x00" * 32)
        assert source._sniff_binary_type(path) == ""

    def test_bom_prefixed_rtf_is_detected(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "bom.bin"
        path.write_bytes(b"\xef\xbb\xbf" + rb"{\rtf1\ansi hi}")
        assert source._sniff_binary_type(path) == ".rtf"

    def test_zip_with_no_known_entries_returns_empty(self, tmp_path):
        import zipfile

        source = _make_source(tmp_path)
        path = tmp_path / "plain.zip"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("unrelated.txt", "hi")
        assert source._sniff_binary_type(path) == ""

    def test_zip_with_docx_entry_detected(self, tmp_path):
        import zipfile

        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("word/document.xml", "<xml/>")
        assert source._sniff_binary_type(path) == ".docx"

    def test_zip_with_odt_entries_detected(self, tmp_path):
        import zipfile

        source = _make_source(tmp_path)
        path = tmp_path / "renamed.bin"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("content.xml", "<xml/>")
            zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        assert source._sniff_binary_type(path) == ".odt"


class TestNormalizeText:
    def test_empty_content_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        assert source._normalize_text("") == ""

    def test_crlf_and_cr_normalized(self, tmp_path):
        source = _make_source(tmp_path)
        assert source._normalize_text("a\r\nb\rc") == "a\nb\nc"


class TestPdfImportErrorAndExceptions:
    def test_pypdf_import_error_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pdf"
        path.write_bytes(b"%PDF-1.4 fake")
        with mock.patch.dict(sys.modules, {"pypdf": None}):
            assert source._extract_pdf_text(path, __import__("time").monotonic()) == ""

    def test_pdf_reader_exception_returns_empty(self, tmp_path):
        pytest.importorskip("pypdf")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pdf"
        path.write_bytes(b"not a real pdf")
        text = source._extract_pdf_text(path, __import__("time").monotonic())
        assert text == ""

    def test_pdf_page_extraction_times_out(self, tmp_path):
        pypdf = pytest.importorskip("pypdf")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pdf"
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        with path.open("wb") as f:
            writer.write(f)

        with mock.patch.object(source, "_timed_out", return_value=True):
            text = source._extract_pdf_text(path, __import__("time").monotonic())
        assert text == ""

    def test_pdf_page_with_extractable_text_is_collected(self, tmp_path):
        pytest.importorskip("pypdf")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pdf"

        fake_page = mock.Mock()
        fake_page.extract_text.return_value = "Real extracted PDF text"
        fake_reader = mock.Mock()
        fake_reader.is_encrypted = False
        fake_reader.pages = [fake_page]

        with mock.patch("pypdf.PdfReader", return_value=fake_reader):
            text = source._extract_pdf_text(path, __import__("time").monotonic())
        assert "Real extracted PDF text" in text

    def test_pdf_page_limit_enforced(self, tmp_path):
        pypdf = pytest.importorskip("pypdf")
        source = _make_source(tmp_path)
        source.MAX_PDF_PAGES = 1
        path = tmp_path / "sample.pdf"
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        with path.open("wb") as f:
            writer.write(f)

        text = source._extract_pdf_text(path, __import__("time").monotonic())
        assert isinstance(text, str)


class TestDocxImportErrorAndExceptions:
    def test_python_docx_import_error_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.docx"
        path.write_bytes(b"PK\x03\x04 fake")
        with mock.patch.dict(sys.modules, {"docx": None}):
            assert source._extract_docx_text(path, __import__("time").monotonic()) == ""

    def test_docx_open_exception_returns_empty(self, tmp_path):
        pytest.importorskip("docx")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.docx"
        path.write_bytes(b"not a real docx")
        assert source._extract_docx_text(path, __import__("time").monotonic()) == ""

    def test_docx_paragraph_extraction_times_out(self, tmp_path):
        docx = pytest.importorskip("docx")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.docx"
        doc = docx.Document()
        doc.add_paragraph("first paragraph")
        doc.add_paragraph("second paragraph")
        doc.save(str(path))

        with mock.patch.object(source, "_timed_out", return_value=True):
            text = source._extract_docx_text(path, __import__("time").monotonic())
        assert text == ""

    def test_docx_table_extraction_times_out(self, tmp_path):
        docx = pytest.importorskip("docx")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.docx"
        doc = docx.Document()
        table1 = doc.add_table(rows=1, cols=1)
        table1.rows[0].cells[0].text = "first table"
        table2 = doc.add_table(rows=1, cols=1)
        table2.rows[0].cells[0].text = "second table"
        doc.save(str(path))

        call_count = {"n": 0}

        def fake_timed_out(started_at):
            call_count["n"] += 1
            # No paragraphs -> first call is for table1 (False), second call
            # (table2) -> True, triggering the tables-loop timeout break.
            return call_count["n"] > 1

        with mock.patch.object(source, "_timed_out", side_effect=fake_timed_out):
            text = source._extract_docx_text(path, __import__("time").monotonic())
        assert "first table" in text
        assert "second table" not in text


class TestXlsxImportErrorAndExceptions:
    def test_openpyxl_import_error_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xlsx"
        path.write_bytes(b"PK\x03\x04 fake")
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            assert source._extract_xlsx_text(path, __import__("time").monotonic()) == ""

    def test_xlsx_open_exception_returns_empty(self, tmp_path):
        pytest.importorskip("openpyxl")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xlsx"
        path.write_bytes(b"not a real xlsx")
        assert source._extract_xlsx_text(path, __import__("time").monotonic()) == ""

    def test_xlsx_sheet_loop_times_out(self, tmp_path):
        openpyxl = pytest.importorskip("openpyxl")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.active["A1"] = "value"
        wb.create_sheet("Sheet2")
        wb.save(str(path))
        wb.close()

        with mock.patch.object(source, "_timed_out", return_value=True):
            text = source._extract_xlsx_text(path, __import__("time").monotonic())
        assert text == ""

    def test_xlsx_row_loop_times_out(self, tmp_path):
        openpyxl = pytest.importorskip("openpyxl")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xlsx"
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet["A1"] = "row1"
        sheet["A2"] = "row2"
        wb.save(str(path))
        wb.close()

        call_count = {"n": 0}

        def fake_timed_out(started_at):
            call_count["n"] += 1
            return call_count["n"] > 1

        with mock.patch.object(source, "_timed_out", side_effect=fake_timed_out):
            text = source._extract_xlsx_text(path, __import__("time").monotonic())
        assert isinstance(text, str)


class TestXlsImportErrorAndExceptions:
    def test_xlrd_import_error_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xls"
        path.write_bytes(b"fake xls bytes")
        with mock.patch.dict(sys.modules, {"xlrd": None}):
            assert source._extract_xls_text(path, __import__("time").monotonic()) == ""

    def test_xls_open_exception_returns_empty(self, tmp_path):
        pytest.importorskip("xlrd")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xls"
        path.write_bytes(b"not a real xls")
        assert source._extract_xls_text(path, __import__("time").monotonic()) == ""

    def test_xls_sheet_loop_times_out(self, tmp_path):
        xlwt = pytest.importorskip("xlwt")
        pytest.importorskip("xlrd")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xls"
        wb = xlwt.Workbook()
        sheet1 = wb.add_sheet("Sheet1")
        sheet1.write(0, 0, "value")
        wb.add_sheet("Sheet2")
        wb.save(str(path))

        with mock.patch.object(source, "_timed_out", return_value=True):
            text = source._extract_xls_text(path, __import__("time").monotonic())
        assert text == ""

    def test_xls_row_loop_times_out(self, tmp_path):
        xlwt = pytest.importorskip("xlwt")
        pytest.importorskip("xlrd")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.xls"
        wb = xlwt.Workbook()
        sheet1 = wb.add_sheet("Sheet1")
        sheet1.write(0, 0, "row1")
        sheet1.write(1, 0, "row2")
        wb.save(str(path))

        call_count = {"n": 0}

        def fake_timed_out(started_at):
            call_count["n"] += 1
            return call_count["n"] > 1

        with mock.patch.object(source, "_timed_out", side_effect=fake_timed_out):
            text = source._extract_xls_text(path, __import__("time").monotonic())
        assert isinstance(text, str)


class TestPptxShapeTextPartsAndImportError:
    def test_pptx_enum_import_error_returns_empty_list(self, tmp_path):
        source = _make_source(tmp_path)
        with mock.patch.dict(sys.modules, {"pptx.enum.shapes": None}):
            assert source._pptx_shape_text_parts([], 0.0) == []

    def test_pptx_import_error_returns_empty_string(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pptx"
        path.write_bytes(b"PK\x03\x04 fake")
        with mock.patch.dict(sys.modules, {"pptx": None}):
            assert source._extract_pptx_text(path, __import__("time").monotonic()) == ""

    def test_pptx_open_exception_returns_empty(self, tmp_path):
        pytest.importorskip("pptx")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pptx"
        path.write_bytes(b"not a real pptx")
        assert source._extract_pptx_text(path, __import__("time").monotonic()) == ""

    def test_pptx_slide_loop_times_out(self, tmp_path):
        pptx = pytest.importorskip("pptx")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pptx"
        prs = pptx.Presentation()
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.save(str(path))

        with mock.patch.object(source, "_timed_out", return_value=True):
            text = source._extract_pptx_text(path, __import__("time").monotonic())
        assert text == ""

    def test_pptx_notes_slide_extracted(self, tmp_path):
        pptx = pytest.importorskip("pptx")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.pptx"
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Title text"
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = "Speaker notes here"
        prs.save(str(path))

        text = source._extract_pptx_text(path, __import__("time").monotonic())
        assert "Speaker notes here" in text

    def test_grouped_shapes_recursed_into(self, tmp_path):
        pytest.importorskip("pptx")
        import time as time_mod

        source = _make_source(tmp_path)
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        inner_shape = mock.Mock()
        inner_shape.shape_type = None
        inner_shape.has_text_frame = True
        inner_shape.text = "nested text"

        group_shape = mock.Mock()
        group_shape.shape_type = MSO_SHAPE_TYPE.GROUP
        group_shape.shapes = [inner_shape]

        parts = source._pptx_shape_text_parts([group_shape], time_mod.monotonic())
        assert parts == ["nested text"]

    def test_table_shape_rows_extracted(self, tmp_path):
        pytest.importorskip("pptx")
        import time as time_mod

        source = _make_source(tmp_path)

        cell1 = mock.Mock(text="A")
        cell2 = mock.Mock(text="B")
        row = mock.Mock(cells=[cell1, cell2])
        table = mock.Mock(rows=[row])

        table_shape = mock.Mock()
        table_shape.shape_type = None
        table_shape.has_text_frame = False
        table_shape.has_table = True
        table_shape.table = table

        parts = source._pptx_shape_text_parts([table_shape], time_mod.monotonic())
        assert parts == ["A | B"]

    def test_shape_level_timeout_breaks_outer_loop(self, tmp_path):
        pytest.importorskip("pptx")
        source = _make_source(tmp_path)

        shape1 = mock.Mock(shape_type=None, has_text_frame=True, text="first shape")
        shape2 = mock.Mock(shape_type=None, has_text_frame=True, text="second shape")

        call_count = {"n": 0}

        def fake_timed_out(started_at):
            call_count["n"] += 1
            return call_count["n"] > 1

        with mock.patch.object(source, "_timed_out", side_effect=fake_timed_out):
            parts = source._pptx_shape_text_parts([shape1, shape2], 0.0)
        assert parts == ["first shape"]

    def test_row_level_timeout_breaks_table_loop(self, tmp_path):
        pytest.importorskip("pptx")
        source = _make_source(tmp_path)

        cell1 = mock.Mock(text="row1")
        row1 = mock.Mock(cells=[cell1])
        cell2 = mock.Mock(text="row2")
        row2 = mock.Mock(cells=[cell2])
        table = mock.Mock(rows=[row1, row2])

        table_shape = mock.Mock(shape_type=None, has_text_frame=False, has_table=True, table=table)

        call_count = {"n": 0}

        def fake_timed_out(started_at):
            call_count["n"] += 1
            # Call 1: outer shape-loop check (False). Call 2: inner row-loop
            # check for row1 (False). Call 3: inner row-loop check for row2
            # (True), breaking the row loop before row2 is processed.
            return call_count["n"] > 2

        with mock.patch.object(source, "_timed_out", side_effect=fake_timed_out):
            parts = source._pptx_shape_text_parts([table_shape], 0.0)
        assert parts == ["row1"]


class TestOdtImportErrorAndExceptions:
    def test_odfpy_import_error_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.odt"
        path.write_bytes(b"PK\x03\x04 fake")
        with mock.patch.dict(sys.modules, {"odf.opendocument": None}):
            assert source._extract_odt_text(path, __import__("time").monotonic()) == ""

    def test_odt_open_exception_returns_empty(self, tmp_path):
        pytest.importorskip("odf.opendocument")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.odt"
        path.write_bytes(b"not a real odt")
        assert source._extract_odt_text(path, __import__("time").monotonic()) == ""

    def test_odt_paragraph_loop_times_out(self, tmp_path):
        odf_text = pytest.importorskip("odf.text")
        odf_doc = pytest.importorskip("odf.opendocument")
        source = _make_source(tmp_path)
        path = tmp_path / "sample.odt"
        doc = odf_doc.OpenDocumentText()
        doc.text.addElement(odf_text.P(text="first"))
        doc.text.addElement(odf_text.P(text="second"))
        doc.save(str(path))

        with mock.patch.object(source, "_timed_out", return_value=True):
            text = source._extract_odt_text(path, __import__("time").monotonic())
        assert text == ""


class TestRtfEdgeCases:
    def test_read_failure_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.rtf"
        path.write_text("{\\rtf1 hi}", encoding="utf-8")
        with mock.patch("builtins.open", side_effect=RuntimeError("io error")):
            assert source._extract_rtf_text(path) == ""

    def test_whitespace_only_content_returns_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "blank.rtf"
        path.write_text("   \n  ", encoding="utf-8")
        assert source._extract_rtf_text(path) == ""

    def test_negative_unicode_escape_decoded(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.rtf"
        # \u-4? decodes to chr(-4 + 65536)
        path.write_text(r"{\rtf1 value: \u-4?}", encoding="utf-8")
        text = source._extract_rtf_text(path)
        assert "value:" in text

    def test_out_of_range_unicode_escape_decodes_to_empty(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.rtf"
        # chr(9999999) raises ValueError inside _decode_unicode's except branch,
        # which should be swallowed and replaced with an empty string.
        path.write_text(r"{\rtf1 before\u9999999? after}", encoding="utf-8")
        text = source._extract_rtf_text(path)
        assert "before" in text
        assert "after" in text

    def test_hex_escape_decoded(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.rtf"
        path.write_text(r"{\rtf1 caf\'e9}", encoding="utf-8")
        text = source._extract_rtf_text(path)
        assert "caf" in text

    def test_single_level_destination_group_removed(self, tmp_path):
        # The destination-removal regex uses `[^{}]*`, which only matches a
        # single (non-nested) brace level -- exercise it with an `\info`
        # destination that has no inner braces.
        source = _make_source(tmp_path)
        path = tmp_path / "sample.rtf"
        path.write_text(
            r"{\rtf1{\info some hidden metadata}\pard visible text\par}",
            encoding="utf-8",
        )
        text = source._extract_rtf_text(path)
        assert "visible text" in text
        assert "hidden metadata" not in text

    def test_parse_exception_falls_back_to_normalized_raw_content(self, tmp_path):
        source = _make_source(tmp_path)
        path = tmp_path / "sample.rtf"
        path.write_text(r"{\rtf1 fallback content}", encoding="utf-8")

        with mock.patch("re.sub", side_effect=RuntimeError("regex boom")):
            text = source._extract_rtf_text(path)
        assert "fallback content" in text


class TestGetTotalCountAndSourceType:
    def test_get_total_count_returns_file_count(self, tmp_path):
        source = _make_source(tmp_path)
        source._files = ["a", "b", "c"]
        assert source.get_total_count() == 3

    def test_source_type_is_unstructured(self, tmp_path):
        source = _make_source(tmp_path)
        assert source.source_type == "unstructured"
